#!/usr/bin/env python3
"""Build the JCCI Time Box Canva-importable master presentation."""

from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/presentations/jcci-time-box/slide-text.md"
OUT = ROOT / "artifacts/jcci-time-box/jcci-time-box-master.pptx"

CREAM = RGBColor(249, 247, 241)
BLACK = RGBColor(20, 21, 23)
GOLD = RGBColor(190, 154, 80)
MUTED = RGBColor(105, 101, 92)
PLACEHOLDER = RGBColor(147, 81, 38)


def read_slides():
    section = None
    slides = []
    for raw in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line.startswith("## "):
            section = line[3:].strip()
        elif line and not line.startswith("#") and section is not None:
            slides.append((section, line))
    if len(slides) != 500:
        raise ValueError(f"Expected 500 slides; found {len(slides)}")
    return slides


def add_text(slide, x, y, w, h, text, size, colour, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.alignment = align
    para.space_after = Pt(0)
    run = para.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    return box


def build():
    content = read_slides()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for idx, (section, line) in enumerate(content, start=1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = CREAM

        add_text(slide, 0.72, 0.32, 10.9, 0.34, "GROWTHCRED  /  JCCI YOUTH COMMISSION", 11, GOLD, True)
        add_text(slide, 0.72, 0.80, 11.9, 0.38, re.sub(r"^\d+\s+—\s+", "", section).upper(), 11, MUTED, True)

        is_placeholder = line.startswith("{") and line.endswith("}")
        size = 31 if len(line) > 115 else 37 if len(line) > 65 else 44
        colour = PLACEHOLDER if is_placeholder else BLACK
        if is_placeholder:
            size = min(size, 31)
        add_text(slide, 0.96, 1.60, 11.4, 4.0, line, size, colour, True, PP_ALIGN.CENTER)

        # The footer is intentionally editable in Canva after import.
        add_text(slide, 0.72, 6.84, 6.0, 0.26, "THE TIME BOX", 11, GOLD, True)
        add_text(slide, 11.75, 6.84, 0.85, 0.26, f"{idx:03d}", 10, MUTED, False, PP_ALIGN.RIGHT)
        if is_placeholder:
            add_text(slide, 4.92, 6.58, 3.5, 0.22, "PRODUCTION PLACEHOLDER", 9, PLACEHOLDER, True, PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Created {len(prs.slides)} slides: {OUT}")


if __name__ == "__main__":
    build()
